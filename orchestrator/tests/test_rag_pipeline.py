"""
Tests for Document Ingestion and MongoDB Hybrid Retrieval RAG Pipeline.
"""

import pytest
import io
import docx
from unittest.mock import AsyncMock, patch, MagicMock

from orchestrator.services.document_ingestion import document_ingestion_service
from orchestrator.services.mongodb_retriever import MongoDBRetriever
from orchestrator.services.mongodb_service import mongodb_service
from orchestrator.config.settings import settings
from orchestrator.models.context_models import RetrievedChunk


@pytest.fixture(autouse=True)
async def init_db_for_test():
    """Force re-initialization of MongoDB service on the active test event loop."""
    mongodb_service._initialized = False
    mongodb_service._client = None
    mongodb_service._db = None
    await mongodb_service.initialize()
    yield


@pytest.fixture
async def clean_db():
    """Ensure document_chunks collection is clean before and after tests."""
    await mongodb_service.document_chunks.delete_many({})
    yield
    await mongodb_service.document_chunks.delete_many({})


class TestDocumentIngestion:
    """Test parsing, cleaning, and adaptive chunking logic."""

    def test_clean_text(self):
        text = "Hello    World! \n\n\n\n New paragraph \t with spaces."
        cleaned = document_ingestion_service.clean_text(text)
        assert cleaned == "Hello World!\n\nNew paragraph with spaces."

    def test_chunk_text_adaptive(self):
        # Generate some text with known word counts
        # target chunk_size = 20 words, overlap = 5 words
        # Make a paragraph with 3 sentences
        sentences = [
            "This is the first sentence of our document.",  # 8 words
            "We are testing adaptive semantic chunking boundaries.",  # 7 words
            "It should group them nicely.",  # 5 words
            "Here is another paragraph containing more details.",  # 7 words
            "This sentence should cross the boundary of the first chunk."  # 10 words
        ]
        text = " ".join(sentences[:3]) + "\n\n" + " ".join(sentences[3:])
        
        with patch.object(settings, "chunk_size", 20), \
             patch.object(settings, "chunk_overlap", 5):
            chunks = document_ingestion_service.chunk_text_adaptive(text)
            
            assert len(chunks) > 0
            # Check fields
            assert "text" in chunks[0]
            assert "chunk_index" in chunks[0]["metadata"]
            assert "paragraph_index" in chunks[0]["metadata"]
            assert "word_count" in chunks[0]["metadata"]

    def test_parse_document_txt(self):
        content = b"Simple text file content."
        parsed = document_ingestion_service.parse_document(content, "test.txt")
        assert parsed == "Simple text file content."

    def test_parse_document_docx(self):
        # Create a mock word doc in memory
        doc = docx.Document()
        doc.add_paragraph("First paragraph.")
        doc.add_paragraph("Second paragraph.")
        
        f_bytes = io.BytesIO()
        doc.save(f_bytes)
        f_bytes.seek(0)
        
        parsed = document_ingestion_service.parse_document(f_bytes.read(), "test.docx")
        assert "First paragraph." in parsed
        assert "Second paragraph." in parsed

    @pytest.mark.asyncio
    async def test_ingest_document(self, clean_db):
        content = b"This is a standalone document. It has several sentences. We will test embedding generation and storage."
        
        mock_embedding = [0.1] * 384
        with patch("orchestrator.services.document_ingestion.embedder_service.embed_batch", new_callable=AsyncMock) as mock_embed:
            mock_embed.return_value = [mock_embedding]
            
            result = await document_ingestion_service.ingest_document(
                file_bytes=content,
                filename="test.txt",
                tenant_id="tenant-123",
                category="manuals"
            )
            
            assert result["status"] == "success"
            assert "doc_id" in result
            assert result["chunks_count"] == 1
            
            # Check DB contents
            db_chunk = await mongodb_service.document_chunks.find_one({"doc_id": result["doc_id"]})
            assert db_chunk is not None
            assert db_chunk["metadata"]["filename"] == "test.txt"
            assert db_chunk["metadata"]["tenant_id"] == "tenant-123"
            assert db_chunk["metadata"]["category"] == "manuals"
            assert db_chunk["embedding"] == mock_embedding


class TestMongoDBRetriever:
    """Test Stage 1 Hybrid Search, Stage 2 Reranking, and Stage 3 Parent-Child Expansion."""

    @pytest.fixture
    async def seed_chunks(self, clean_db):
        doc_id = "test-doc-id-1"
        chunks = [
            {
                "doc_id": doc_id,
                "chunk_id": f"{doc_id}_0",
                "text": "Photosynthesis converts solar energy into chemical energy in green plants.",
                "embedding": [0.9, 0.1, 0.0] + [0.0]*381,
                "metadata": {
                    "filename": "biology.txt",
                    "chunk_index": 0,
                    "tenant_id": "tenant-test",
                    "category": "biology"
                }
            },
            {
                "doc_id": doc_id,
                "chunk_id": f"{doc_id}_1",
                "text": "Chlorophyll is the primary pigment involved in absorbing light energy.",
                "embedding": [0.1, 0.9, 0.0] + [0.0]*381,
                "metadata": {
                    "filename": "biology.txt",
                    "chunk_index": 1,
                    "tenant_id": "tenant-test",
                    "category": "biology"
                }
            },
            {
                "doc_id": doc_id,
                "chunk_id": f"{doc_id}_2",
                "text": "The Calvin cycle takes place in the stroma of chloroplasts.",
                "embedding": [0.0, 0.1, 0.9] + [0.0]*381,
                "metadata": {
                    "filename": "biology.txt",
                    "chunk_index": 2,
                    "tenant_id": "tenant-test",
                    "category": "biology"
                }
            }
        ]
        await mongodb_service.document_chunks.insert_many(chunks)
        return doc_id

    @pytest.mark.asyncio
    async def test_hybrid_and_rerank_search(self, seed_chunks):
        retriever = MongoDBRetriever()
        
        # Mock embedder and reranker to return predictable outputs
        mock_query_vector = [0.9, 0.1, 0.0] + [0.0]*381
        mock_rerank_scores = [0.95, 0.50, 0.10]
        
        with patch("orchestrator.services.mongodb_retriever.embedder_service.embed_query", new_callable=AsyncMock) as mock_embed, \
             patch("orchestrator.services.mongodb_retriever.reranker_service.rerank", return_value=mock_rerank_scores), \
             patch.object(settings, "parent_child_window", 0), \
             patch.object(settings, "skip_rerank_similarity_threshold", 2.0):  # Disable parent-child context expansion for simplified test
            
            mock_embed.return_value = mock_query_vector
            
            results = await retriever.search(
                query="photosynthesis light",
                filters={"tenant_id": "tenant-test"}
            )
            
            assert len(results) > 0
            assert isinstance(results[0], RetrievedChunk)
            assert results[0].source == "biology.txt"
            
            # The highest score (reranked) should be the first one
            assert results[0].id == "test-doc-id-1_0"
            assert results[0].score == 0.95

    @pytest.mark.asyncio
    async def test_parent_child_expansion(self, seed_chunks):
        retriever = MongoDBRetriever()
        
        mock_query_vector = [0.1, 0.9, 0.0] + [0.0]*381
        mock_rerank_scores = [0.99]
        
        with patch("orchestrator.services.mongodb_retriever.embedder_service.embed_query", new_callable=AsyncMock) as mock_embed, \
             patch("orchestrator.services.mongodb_retriever.reranker_service.rerank", return_value=mock_rerank_scores), \
             patch.object(settings, "parent_child_window", 1), \
             patch.object(settings, "rerank_top_k", 1), \
             patch.object(settings, "skip_rerank_similarity_threshold", 2.0): # expand with window=1 (returns original + index-1 + index+1)
            
            mock_embed.return_value = mock_query_vector
            
            results = await retriever.search(
                query="chlorophyll primary pigment",
                filters={"tenant_id": "tenant-test"}
            )
            
            assert len(results) == 1
            # Text should be expanded to include chunk 0, 1, and 2
            text = results[0].text
            assert "Photosynthesis converts solar energy" in text
            assert "Chlorophyll is the primary pigment" in text
            assert "The Calvin cycle takes place" in text
            
            # Metadata should flag expansion
            assert results[0].metadata["expanded"] is True
            assert results[0].metadata["original_chunk_index"] == 1
            assert results[0].metadata["expanded_chunks_count"] == 3

    @pytest.mark.asyncio
    async def test_rerank_bypass(self, seed_chunks):
        retriever = MongoDBRetriever()
        mock_query_vector = [0.9, 0.1, 0.0] + [0.0]*381
        
        with patch("orchestrator.services.mongodb_retriever.embedder_service.embed_query", new_callable=AsyncMock) as mock_embed, \
             patch("orchestrator.services.mongodb_retriever.reranker_service.rerank") as mock_rerank, \
             patch.object(settings, "parent_child_window", 0), \
             patch.object(settings, "skip_rerank_similarity_threshold", 0.85):
            
            mock_embed.return_value = mock_query_vector
            
            results = await retriever.search(
                query="photosynthesis light",
                filters={"tenant_id": "tenant-test"}
            )
            
            assert len(results) > 0
            mock_rerank.assert_not_called()
            assert results[0].score > 0.0

    @pytest.mark.asyncio
    async def test_early_abort_db_miss(self, seed_chunks):
        retriever = MongoDBRetriever()
        mock_query_vector = [0.0, 0.0, 0.0] + [0.0]*381
        
        with patch("orchestrator.services.mongodb_retriever.embedder_service.embed_query", new_callable=AsyncMock) as mock_embed, \
             patch.object(settings, "early_abort_similarity_threshold", 0.35):
            
            mock_embed.return_value = mock_query_vector
            
            results = await retriever.search(
                query="something totally unrelated",
                filters={"tenant_id": "tenant-test"}
            )
            
            assert len(results) == 1
            assert results[0].metadata.get("db_miss") is True
            assert results[0].id == "db_miss_chunk"
